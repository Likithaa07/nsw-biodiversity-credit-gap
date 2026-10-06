# STEP 1: Imports and file locations
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
PROCESSED = ROOT / "data" / "processed"
OUTPUTS = ROOT / "outputs"
OUTPUTS.mkdir(exist_ok=True)
EXCEL = OUTPUTS / "powerbi_data.xlsx"


# STEP 2: Load the results from the earlier scripts
gap = pd.read_csv(PROCESSED / "gap_ranking.csv")
price_year = pd.read_csv(PROCESSED / "price_by_year.csv")
price_scarcity = pd.read_csv(PROCESSED / "price_by_scarcity.csv")
quality = pd.read_csv(PROCESSED / "data_quality_log.csv")
supply = pd.read_csv(PROCESSED / "supply_clean.csv")
transactions = pd.read_csv(PROCESSED / "transactions_clean.csv", parse_dates=["transaction_date"])


# STEP 3: Market overview tables
supply_summary = (supply.groupby(["credit_status", "credit_class", "ibra_region"], dropna=False)
                  ["number_of_credits"].sum().reset_index()
                  .rename(columns={"number_of_credits": "credits"}))
supply_summary["ibra_region"] = supply_summary["ibra_region"].fillna("Not recorded")

activity = transactions[transactions["transaction_type"].isin(["Transfer", "Retire"])].copy()
activity["year"] = activity["transaction_date"].dt.year
activity_by_year = (activity.groupby(["year", "transaction_type", "credit_class"])
                    ["number_of_credits"].sum().reset_index()
                    .rename(columns={"number_of_credits": "credits"}))


# STEP 4: Write every table to its own sheet
sheets = {
    "gap_ranking": gap,
    "price_by_year": price_year,
    "price_by_scarcity": price_scarcity,
    "supply_summary": supply_summary,
    "activity_by_year": activity_by_year,
    "data_quality_log": quality,
}
with pd.ExcelWriter(EXCEL, engine="openpyxl") as writer:
    for name, df in sheets.items():
        df.to_excel(writer, sheet_name=name, index=False)


# STEP 5: Format each sheet as an Excel table (Power BI in the browser looks for tables)
wb = load_workbook(EXCEL)
for name, df in sheets.items():
    ws = wb[name]
    ref = f"A1:{get_column_letter(df.shape[1])}{df.shape[0] + 1}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
wb.save(EXCEL)

for name, df in sheets.items():
    print(f"{name:<18} {len(df):>5} rows")
print(f"\nSaved {EXCEL}")
